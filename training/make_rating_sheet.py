"""Turns eval's samples.csv into a clean, BLIND rating sheet (Excel) for human raters.

  python training/make_rating_sheet.py --csv training/out/eval/samples.csv --out training/out/eval/rating_sheet

Creates  rating_sheet.xlsx  (give this to raters: systems are hidden as A/B/C, order shuffled per case)
and      rating_sheet_KEY.csv (keep this: which letter was which system). Needs: pip install openpyxl
"""
import argparse
import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

BASE = ["id", "test", "value", "status"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="training/out/eval/samples.csv")
    ap.add_argument("--out", default="training/out/eval/rating_sheet")
    ap.add_argument("--skip", nargs="*", default=["qwen3:4b"], help="systems to leave out (e.g. the thinking model)")
    ap.add_argument("--seed", type=int, default=7)
    a = ap.parse_args()

    rows = list(csv.DictReader(open(a.csv, encoding="utf-8-sig")))
    systems = [k for k in rows[0] if k not in BASE and not k.startswith("rating_") and k not in a.skip]
    letters = [chr(65 + i) for i in range(len(systems))]
    rng = random.Random(a.seed)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rating"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(["Instructions: for each case read the 3 explanations (A, B, C) and give each a score 1-5 in the yellow cells. "
               "1 = confusing / unnatural Bangla, 3 = understandable but stiff, 5 = clear and natural, like a person explaining. "
               "Rate the Bangla wording only."])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + 2 * len(systems))
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True)
    ws.row_dimensions[1].height = 48
    head = ["Case", "Test (value, status)"]
    for L in letters:
        head += [f"Explanation {L}", f"Score {L} (1-5)"]
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = border

    key = [["case"] + [f"letter_{L}" for L in letters]]
    for i, r in enumerate(rows, 1):
        order = systems[:]
        rng.shuffle(order)
        line = [i, f"{r['test']} ({float(r['value']):g}, {r['status']})"]
        for s in order:
            line += [r[s].strip(), None]
        ws.append(line)
        key.append([i] + order)
    dv = DataValidation(type="whole", operator="between", formula1=1, formula2=5, allow_blank=True,
                        showErrorMessage=True, errorTitle="Score", error="Enter a whole number from 1 to 5")
    ws.add_data_validation(dv)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        for j in range(2, len(row)):
            c = row[j]
            c.border = border
            if j % 2 == 0:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.fill, c.alignment = yellow, Alignment(horizontal="center", vertical="center")
                dv.add(c)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    for j in range(len(systems)):
        ws.column_dimensions[chr(67 + 2 * j)].width = 55
        ws.column_dimensions[chr(68 + 2 * j)].width = 10
    ws.freeze_panes = "C3"

    avg = wb.create_sheet("Average")
    avg.append(["Letter", "Average score"])
    for j, L in enumerate(letters):
        col = chr(68 + 2 * j)
        avg.append([L, f"=IFERROR(AVERAGE(Rating!{col}3:{col}{ws.max_row}),\"\")"])
    avg.column_dimensions["A"].width = 10
    avg.column_dimensions["B"].width = 16

    out = Path(a.out)
    wb.save(out.with_suffix(".xlsx"))
    with open(str(out) + "_KEY.csv", "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(key)
    print("wrote", out.with_suffix(".xlsx"), "and", str(out) + "_KEY.csv", "| systems:", systems)


if __name__ == "__main__":
    main()
